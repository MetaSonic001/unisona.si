"""Turn raw sources (URLs, websites, files, text) into clean markdown documents."""
from __future__ import annotations

import asyncio
import io
import re
from dataclasses import dataclass, field
from urllib.parse import urljoin, urldefrag, urlparse

import httpx

from ..log import get_logger

log = get_logger("extract")
UA = "Mozilla/5.0 (compatible; UnisonaBot/1.0; +https://unisona.si)"


@dataclass
class Doc:
    title: str
    text: str
    url: str | None = None
    meta: dict = field(default_factory=dict)


@dataclass
class TableDoc:
    name: str
    rows: list[dict]
    columns: list[str]


def _clean(text: str) -> str:
    text = re.sub(r"\n{3,}", "\n\n", text or "")
    return text.strip()


async def fetch_html(url: str, client: httpx.AsyncClient) -> tuple[str, str]:
    r = await client.get(url, follow_redirects=True)
    r.raise_for_status()
    ctype = r.headers.get("content-type", "")
    if "pdf" in ctype:
        return "__PDF__", r.content.hex()
    return str(r.url), r.text


def html_to_doc(url: str, html: str) -> Doc:
    import trafilatura

    md = trafilatura.extract(html, url=url, output_format="markdown", include_tables=True, include_links=False,
                             favor_recall=True) or ""
    title = ""
    meta = trafilatura.extract_metadata(html)
    if meta and meta.title:
        title = meta.title
    if len(md) < 200:  # fallback: visible text via BeautifulSoup
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(html, "html.parser")
        for t in soup(["script", "style", "noscript", "svg", "nav", "footer"]):
            t.decompose()
        title = title or (soup.title.string.strip() if soup.title and soup.title.string else "")
        md = soup.get_text("\n", strip=True)
    return Doc(title=title or urlparse(url).path or url, text=_clean(md), url=url)


async def extract_url(url: str) -> list[Doc]:
    async with httpx.AsyncClient(timeout=25, headers={"User-Agent": UA}) as c:
        final, html = await fetch_html(url, c)
        if final == "__PDF__":
            return extract_pdf(bytes.fromhex(html), url)
        return [html_to_doc(final, html)]


def _same_site(base: str, url: str) -> bool:
    a, b = urlparse(base), urlparse(url)
    return a.netloc.replace("www.", "") == b.netloc.replace("www.", "")


_SKIP = re.compile(r"\.(png|jpe?g|gif|svg|webp|ico|css|js|zip|mp4|mp3|woff2?)$|/(cart|checkout|login|signin|account|wp-admin)", re.I)
_PRIORITY = re.compile(r"(about|pricing|price|plans|faq|help|support|contact|service|product|feature|policy|terms|refund|shipping|return|team|hours|location|menu|course)", re.I)


async def _sitemap_urls(root: str, client: httpx.AsyncClient) -> list[str]:
    urls: list[str] = []
    for path in ("/sitemap.xml", "/sitemap_index.xml"):
        try:
            r = await client.get(urljoin(root, path), follow_redirects=True)
            if r.status_code != 200 or "<" not in r.text:
                continue
            locs = re.findall(r"<loc>\s*([^<\s]+)\s*</loc>", r.text)
            for loc in locs:
                if loc.endswith(".xml") and len(urls) < 400:
                    try:
                        sub = await client.get(loc, follow_redirects=True)
                        urls += re.findall(r"<loc>\s*([^<\s]+)\s*</loc>", sub.text)
                    except Exception:
                        pass
                else:
                    urls.append(loc)
            if urls:
                break
        except Exception:
            continue
    return urls


async def crawl_website(start_url: str, max_pages: int = 20, on_progress=None) -> list[Doc]:
    """Sitemap first (prioritising about/pricing/faq/contact pages), then breadth-first links."""
    if not start_url.startswith("http"):
        start_url = "https://" + start_url
    docs: list[Doc] = []
    seen: set[str] = set()
    async with httpx.AsyncClient(timeout=20, headers={"User-Agent": UA}) as c:
        queue: list[str] = [start_url]
        sm = [u for u in await _sitemap_urls(start_url, c) if _same_site(start_url, u) and not _SKIP.search(u)]
        sm.sort(key=lambda u: (0 if _PRIORITY.search(u) else 1, len(u)))
        queue += sm[: max_pages * 2]
        sem = asyncio.Semaphore(5)

        async def visit(url: str) -> list[str]:
            async with sem:
                try:
                    final, html = await fetch_html(url, c)
                    if final == "__PDF__":
                        docs.extend(extract_pdf(bytes.fromhex(html), url))
                        return []
                    doc = html_to_doc(final, html)
                    if len(doc.text) > 120:
                        docs.append(doc)
                        if on_progress:
                            await on_progress(len(docs), url)
                    links = re.findall(r'href=["\']([^"\'#]+)', html)
                    return [urldefrag(urljoin(final, h))[0] for h in links]
                except Exception as e:
                    log.debug(f"crawl skip {url}: {e}")
                    return []

        while queue and len(docs) < max_pages:
            batch = []
            while queue and len(batch) < 5:
                u = queue.pop(0).rstrip("/")
                if u in seen or not _same_site(start_url, u) or _SKIP.search(u):
                    continue
                seen.add(u)
                batch.append(u)
            if not batch:
                continue
            results = await asyncio.gather(*(visit(u) for u in batch))
            new = [u for links in results for u in links if u.rstrip("/") not in seen and _same_site(start_url, u) and not _SKIP.search(u)]
            new.sort(key=lambda u: 0 if _PRIORITY.search(u) else 1)
            queue.extend(new[:60])
    log.info(f"Crawled {start_url}: {len(docs)} pages")
    return docs[:max_pages]


def extract_pdf(data: bytes, name: str) -> list[Doc]:
    import pymupdf
    import pymupdf4llm

    doc = pymupdf.open(stream=data, filetype="pdf")
    pages = pymupdf4llm.to_markdown(doc, page_chunks=True, show_progress=False)
    out = []
    for p in pages:
        text = _clean(p.get("text", ""))
        if len(text) > 30:
            page_no = (p.get("metadata") or {}).get("page", len(out) + 1)
            out.append(Doc(title=f"{name} · p.{page_no}", text=text, meta={"page": page_no}))
    if not out:
        log.warning(f"{name}: PDF has no extractable text (scanned?). Consider OCR.")
    return out


def extract_docx(data: bytes, name: str) -> list[Doc]:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = []
    for p in d.paragraphs:
        if not p.text.strip():
            continue
        style = (p.style.name or "").lower() if p.style else ""
        if style.startswith("heading"):
            level = int(re.sub(r"\D", "", style) or 2)
            parts.append("#" * min(level, 4) + " " + p.text.strip())
        else:
            parts.append(p.text.strip())
    for t in d.tables:
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
        parts.append("\n".join(rows))
    return [Doc(title=name, text=_clean("\n\n".join(parts)))]


def extract_pptx(data: bytes, name: str) -> list[Doc]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
        if texts:
            out.append(Doc(title=f"{name} · slide {i}", text=_clean("\n".join(texts)), meta={"slide": i}))
    return out


def extract_table(data: bytes, name: str) -> TableDoc:
    import csv

    lower = name.lower()
    if lower.endswith((".xlsx", ".xlsm", ".xls")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(next(rows_iter))]
        rows = [{header[i]: ("" if v is None else v) for i, v in enumerate(r[: len(header)])} for r in rows_iter if any(v is not None for v in r)]
    else:
        text = data.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        header = [h.strip() for h in (reader.fieldnames or [])]
        rows = [{(k or "").strip(): v for k, v in r.items()} for r in reader]
    return TableDoc(name=name, rows=rows, columns=header)


def extract_file(data: bytes, filename: str) -> list[Doc] | TableDoc:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_pdf(data, filename)
    if lower.endswith(".docx"):
        return extract_docx(data, filename)
    if lower.endswith(".pptx"):
        return extract_pptx(data, filename)
    if lower.endswith((".csv", ".xlsx", ".xlsm", ".xls", ".tsv")):
        return extract_table(data, filename)
    if lower.endswith((".html", ".htm")):
        return [html_to_doc(filename, data.decode("utf-8", errors="replace"))]
    return [Doc(title=filename, text=_clean(data.decode("utf-8", errors="replace")))]
